import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE = "mobiles.xlsx"


# ---------- Excel ----------
def create_excel():
    if not os.path.exists(FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Mobiles"

        ws.append([
            "Mobile ID",
            "Brand",
            "Model",
            "Price",
            "Quantity"
        ])

        wb.save(FILE)


def add_mobile():
    mobile_id = id_entry.get()
    brand = brand_entry.get()
    model = model_entry.get()
    price = price_entry.get()
    quantity = quantity_entry.get()

    if mobile_id == "" or brand == "" or model == "":
        messagebox.showwarning("Warning", "Please fill all fields")
        return

    wb = load_workbook(FILE)
    ws = wb.active

    ws.append([
        mobile_id,
        brand,
        model,
        price,
        quantity
    ])

    wb.save(FILE)

    messagebox.showinfo("Success", "Mobile added successfully")

    clear_fields()
    show_mobiles()


def show_mobiles():
    for item in table.get_children():
        table.delete(item)

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        table.insert("", tk.END, values=row)


def clear_fields():
    id_entry.delete(0, tk.END)
    brand_entry.delete(0, tk.END)
    model_entry.delete(0, tk.END)
    price_entry.delete(0, tk.END)
    quantity_entry.delete(0, tk.END)


def delete_mobile():
    selected = table.selection()

    if not selected:
        messagebox.showwarning("Warning", "Select a mobile first")
        return

    item = table.item(selected[0])
    mobile_id = item["values"][0]

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == mobile_id:
            ws.delete_rows(row[0].row)
            break

    wb.save(FILE)

    messagebox.showinfo("Success", "Mobile deleted successfully")

    show_mobiles()


def search_mobile():
    search = search_entry.get().lower()

    for item in table.get_children():
        table.delete(item)

    wb = load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        if (search in str(row[0]).lower()
                or search in str(row[1]).lower()
                or search in str(row[2]).lower()):

            table.insert("", tk.END, values=row)


def logout():
    root.destroy()


# ---------- Login ----------
def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "pratik" and password == "pratik123":
        login_frame.pack_forget()
        main_frame.pack(fill="both", expand=True)
        show_mobiles()
    else:
        messagebox.showerror("Error", "Invalid Username or Password")


# ---------- Main Window ----------
create_excel()

root = tk.Tk()
root.title("Mobile Shop Management System")
root.geometry("500x650")

# ---------- Login Frame ----------
login_frame = tk.Frame(root)

tk.Label(
    login_frame,
    text="MOBILE SHOP",
    font=("Arial", 22, "bold")
).pack(pady=20)

tk.Label(
    login_frame,
    text="Username"
).pack()

username_entry = tk.Entry(login_frame)
username_entry.pack(pady=5)

tk.Label(
    login_frame,
    text="Password"
).pack()

password_entry = tk.Entry(
    login_frame,
    show="*"
)
password_entry.pack(pady=5)

tk.Button(
    login_frame,
    text="LOGIN",
    command=login,
    width=15
).pack(pady=20)

login_frame.pack(
    fill="both",
    expand=True
)


# ---------- Main Frame ----------
main_frame = tk.Frame(root)

tk.Label(
    main_frame,
    text="Mobile Shop Management System",
    font=("Arial", 18, "bold")
).pack(pady=10)


# ---------- Add Mobile ----------
form_frame = tk.Frame(main_frame)
form_frame.pack(pady=5)

tk.Label(form_frame, text="Mobile ID").grid(
    row=0, column=0, padx=5, pady=5
)

id_entry = tk.Entry(form_frame)
id_entry.grid(row=0, column=1)


tk.Label(form_frame, text="Brand").grid(
    row=1, column=0, padx=5, pady=5
)

brand_entry = tk.Entry(form_frame)
brand_entry.grid(row=1, column=1)


tk.Label(form_frame, text="Model").grid(
    row=2, column=0, padx=5, pady=5
)

model_entry = tk.Entry(form_frame)
model_entry.grid(row=2, column=1)


tk.Label(form_frame, text="Price").grid(
    row=3, column=0, padx=5, pady=5
)

price_entry = tk.Entry(form_frame)
price_entry.grid(row=3, column=1)


tk.Label(form_frame, text="Quantity").grid(
    row=4, column=0, padx=5, pady=5
)

quantity_entry = tk.Entry(form_frame)
quantity_entry.grid(row=4, column=1)


# ---------- Buttons ----------
button_frame = tk.Frame(main_frame)
button_frame.pack(pady=10)

tk.Button(
    button_frame,
    text="Add Mobile",
    command=add_mobile
).grid(row=0, column=0, padx=5)

tk.Button(
    button_frame,
    text="Delete",
    command=delete_mobile
).grid(row=0, column=1, padx=5)

tk.Button(
    button_frame,
    text="Clear",
    command=clear_fields
).grid(row=0, column=2, padx=5)

tk.Button(
    button_frame,
    text="Logout",
    command=logout
).grid(row=0, column=3, padx=5)


# ---------- Search ----------
search_frame = tk.Frame(main_frame)
search_frame.pack(pady=5)

tk.Label(
    search_frame,
    text="Search"
).pack(side=tk.LEFT)

search_entry = tk.Entry(search_frame)
search_entry.pack(side=tk.LEFT, padx=5)

tk.Button(
    search_frame,
    text="Search",
    command=search_mobile
).pack(side=tk.LEFT)


# ---------- Table ----------
table_frame = tk.Frame(main_frame)
table_frame.pack(
    fill="both",
    expand=True,
    padx=10,
    pady=10
)

columns = (
    "ID",
    "Brand",
    "Model",
    "Price",
    "Quantity"
)

table = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings"
)

for col in columns:
    table.heading(col, text=col)
    table.column(col, width=80)

table.pack(
    fill="both",
    expand=True
)


root.mainloop()